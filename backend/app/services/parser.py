"""文档解析：把上传的文件转成「带页码的文本段落」。

所有函数都从内存字节解析，不落临时文件 —— 上传流本来就在内存里，
少一次磁盘往返，也避免临时文件残留。

页码说明：
  PDF / PPTX 能精确到页；DOCX / XLSX / 纯文本没有页的概念，page=None。
  页码会被切片器继承，最终写进片段，用于前端溯源定位。

不支持的格式会抛出 ParseError 并给出可操作的提示，
而不是静默产出空文本 —— 空文档在问答侧表现为「搜不到」，
排查起来非常费劲。
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
from dataclasses import dataclass, field
from html import unescape
from typing import Any, Dict, List, Optional

from app.services.textutil import clean_text

logger = logging.getLogger(__name__)


class ParseError(RuntimeError):
    pass


@dataclass
class Section:
    """一段连续文本，尽量带上页码。"""

    text: str
    page: Optional[int] = None


@dataclass
class ParsedDocument:
    ext: str
    sections: List[Section] = field(default_factory=list)
    title: Optional[str] = None
    meta: Dict[str, Any] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.sections if s.text)

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.sections)

    @property
    def is_empty(self) -> bool:
        return self.char_count == 0


# ---------------------------------------------------------------- 字节解码

_ENCODINGS = ("utf-8", "gb18030", "utf-16", "big5", "latin-1")


def decode_text(content: bytes) -> str:
    for enc in _ENCODINGS:
        try:
            return content.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("utf-8", errors="replace")


# ---------------------------------------------------------------- HTML

_SCRIPT_RE = re.compile(r"<(script|style)\b[^>]*>.*?</\1>", re.S | re.I)
_BLOCK_RE = re.compile(r"</?(p|div|br|li|tr|h[1-6]|section|article)\b[^>]*>", re.I)
_TAG_RE = re.compile(r"<[^>]+>")


def html_to_text(raw: str) -> str:
    raw = _SCRIPT_RE.sub(" ", raw)
    raw = _BLOCK_RE.sub("\n", raw)
    raw = _TAG_RE.sub(" ", raw)
    return clean_text(unescape(raw))


# ---------------------------------------------------------------- 各格式

def _parse_pdf(content: bytes) -> ParsedDocument:
    try:
        import pymupdf as fitz  # PyMuPDF >= 1.24 的新模块名
    except ImportError:
        try:
            import fitz  # 旧的模块名，1.24 起已标记废弃
        except ImportError as exc:  # pragma: no cover
            raise ParseError("服务器缺少 PyMuPDF 依赖，无法解析 PDF") from exc

    sections: List[Section] = []
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PDF 打开失败（文件可能已损坏或被加密）：{exc}") from exc

    try:
        for idx, page in enumerate(doc, start=1):
            try:
                text = page.get_text("text")
            except Exception as exc:  # noqa: BLE001
                logger.warning("PDF 第 %d 页文本提取失败：%s", idx, exc)
                continue
            text = clean_text(text or "")
            if text:
                sections.append(Section(text=text, page=idx))
        page_count = doc.page_count
        meta = dict(doc.metadata or {})
    finally:
        doc.close()

    if not sections:
        raise ParseError(
            "PDF 中没有可提取的文字。若这是扫描件（图片型 PDF），"
            "需要先做 OCR 再上传 —— 当前内核不含 OCR 能力。"
        )

    title = clean_text(str(meta.get("title") or "")) or None
    return ParsedDocument(
        ext="pdf", sections=sections, title=title,
        meta={"page_count": page_count, "pdf_meta": {k: v for k, v in meta.items() if v}},
    )


def _parse_docx(content: bytes) -> ParsedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ParseError("服务器缺少 python-docx 依赖，无法解析 DOCX") from exc

    try:
        doc = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"DOCX 打开失败：{exc}") from exc

    parts: List[str] = []

    def _iter_block_items(document):
        """按文档真实顺序遍历段落与表格（python-docx 默认会拆散顺序）。"""
        from docx.document import Document as _Doc
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table, _Cell
        from docx.text.paragraph import Paragraph

        def _walk(parent):
            if isinstance(parent, _Cell):
                parent_elm = parent._tc
            elif isinstance(parent, _Doc):
                parent_elm = parent.element.body
            else:  # pragma: no cover
                return
            for child in parent_elm.iterchildren():
                if isinstance(child, CT_P):
                    yield Paragraph(child, parent)
                elif isinstance(child, CT_Tbl):
                    yield Table(child, parent)

        yield from _walk(document)

    table_count = 0
    try:
        for block in _iter_block_items(doc):
            if hasattr(block, "rows"):
                table_count += 1
                rows: List[str] = []
                for row in block.rows:
                    cells = [clean_text(c.text or "") for c in row.cells]
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append("\n".join(rows))
            else:
                text = clean_text(block.text or "")
                if text:
                    style = getattr(getattr(block, "style", None), "name", "") or ""
                    if style.lower().startswith("heading"):
                        parts.append(f"\n{text}\n")
                    else:
                        parts.append(text)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"DOCX 内容读取失败：{exc}") from exc

    text = clean_text("\n\n".join(parts))
    if not text:
        raise ParseError("DOCX 中没有可提取的文字内容。")

    return ParsedDocument(
        ext="docx",
        sections=[Section(text=text, page=None)],
        meta={"tables": table_count},
    )


def _parse_pptx(content: bytes) -> ParsedDocument:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise ParseError("服务器缺少 python-pptx 依赖，无法解析 PPTX") from exc

    try:
        prs = Presentation(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"PPTX 打开失败：{exc}") from exc

    sections: List[Section] = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts: List[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for para in shape.text_frame.paragraphs:
                    line = clean_text("".join(r.text or "" for r in para.runs))
                    if line:
                        parts.append(line)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [clean_text(c.text or "") for c in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
        text = clean_text("\n".join(parts))
        if text:
            sections.append(Section(text=text, page=idx))

    if not sections:
        raise ParseError("PPTX 中没有可提取的文字内容。")

    return ParsedDocument(ext="pptx", sections=sections, meta={"slides": len(prs.slides)})


def _parse_xlsx(content: bytes) -> ParsedDocument:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ParseError("服务器缺少 openpyxl 依赖，无法解析 XLSX") from exc

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"XLSX 打开失败：{exc}") from exc

    sections: List[Section] = []
    try:
        for sheet in wb.worksheets:
            lines: List[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    lines.append(" | ".join(cells))
            text = clean_text("\n".join(lines))
            if text:
                sections.append(Section(text=text, page=None))
    finally:
        wb.close()

    if not sections:
        raise ParseError("XLSX 中没有可提取的单元格内容。")

    return ParsedDocument(ext="xlsx", sections=sections, meta={"sheets": len(wb.sheetnames)})


def _parse_text(content: bytes, ext: str) -> ParsedDocument:
    raw = decode_text(content)

    if ext in ("html", "htm"):
        text = html_to_text(raw)
        sections = [Section(text=text, page=None)] if text else []
    elif ext == "json":
        try:
            data = json.loads(raw)
            text = clean_text(json.dumps(data, ensure_ascii=False, indent=2))
        except (ValueError, TypeError):
            text = clean_text(raw)
        sections = [Section(text=text, page=None)] if text else []
    elif ext == "csv":
        rows: List[str] = []
        try:
            reader = csv.reader(io.StringIO(raw))
            for row in reader:
                cells = [c.strip() for c in row]
                if any(cells):
                    rows.append(" | ".join(cells))
            text = clean_text("\n".join(rows))
        except Exception:  # noqa: BLE001
            text = clean_text(raw)
        sections = [Section(text=text, page=None)] if text else []
    else:
        text = clean_text(raw)
        sections = [Section(text=text, page=None)] if text else []

    if not sections:
        raise ParseError("文件内容为空，或全部是空白字符。")

    return ParsedDocument(ext=ext, sections=sections)


_UNSUPPORTED = {
    "doc": "旧版 .doc 二进制格式。请另存为 .docx 或导出 PDF 后上传。",
    "ppt": "旧版 .ppt 二进制格式。请另存为 .pptx 或导出 PDF 后上传。",
    "xls": "旧版 .xls 二进制格式。请另存为 .xlsx 或导出 CSV 后上传。",
    "epub": "EPUB 电子书。请先转换为 PDF 或纯文本后上传。",
}

_TEXT_EXT = {"txt", "md", "markdown", "csv", "json", "log", "html", "htm", "text"}


def parse_bytes(content: bytes, filename: str, ext: Optional[str] = None) -> ParsedDocument:
    if not content:
        raise ParseError("文件内容为空。")

    ext = (ext or filename.rsplit(".", 1)[-1] if "." in filename else "").lower().lstrip(".")

    if ext == "pdf":
        doc = _parse_pdf(content)
    elif ext == "docx":
        doc = _parse_docx(content)
    elif ext == "pptx":
        doc = _parse_pptx(content)
    elif ext == "xlsx":
        doc = _parse_xlsx(content)
    elif ext in _TEXT_EXT:
        doc = _parse_text(content, ext)
    elif ext in _UNSUPPORTED:
        raise ParseError(_UNSUPPORTED[ext])
    else:
        supported = "pdf docx pptx xlsx txt md csv json html"
        raise ParseError(f"暂不支持 .{ext} 格式的解析。当前支持：{supported}")

    if doc.is_empty:
        raise ParseError("解析结果为空，文件可能没有可检索的文字内容。")
    return doc


def parse_path(path: str, ext: Optional[str] = None) -> ParsedDocument:
    from pathlib import Path as _Path

    p = _Path(path)
    if not p.exists():
        raise ParseError(f"文件不存在：{path}")
    return parse_bytes(p.read_bytes(), p.name, ext)
